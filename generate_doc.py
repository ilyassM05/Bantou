import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def add_heading(doc, text, level=1):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0, 51, 102) # Dark Blue
        
def add_table_data(doc, data, headers):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Light Shading Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Add headers
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        for paragraph in hdr_cells[i].paragraphs:
            for run in paragraph.runs:
                run.bold = True
                
    # Add rows
    for row_data in data:
        row_cells = table.add_row().cells
        for i, cell_data in enumerate(row_data):
            row_cells[i].text = str(cell_data)
            
    doc.add_paragraph("") # Space after table

def create_document():
    doc = Document()
    
    # Cover Page
    doc.add_paragraph('\n'*10)
    title = doc.add_paragraph('Fiche Technique / Technical Specification')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.size = Pt(28)
        run.bold = True
        run.font.color.rgb = RGBColor(0, 51, 102)
        
    subtitle = doc.add_paragraph('Bantou - Community Management Application')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in subtitle.runs:
        run.font.size = Pt(18)
        run.italic = True
        
    doc.add_page_break()

    # 1. Project Overview
    add_heading(doc, '1. Project Overview')
    overview_data = [
        ["Project Name", "Bantou"],
        ["Type", "Mobile Application (MVP)"],
        ["Platform", "Android & iOS (Flutter)"],
        ["Purpose", "Community and Association Management"],
        ["Architecture Pattern", "Client-Server Architecture with Real-time Socket Communication"]
    ]
    add_table_data(doc, overview_data, ["Attribute", "Description"])

    # 2. Project Objectives
    add_heading(doc, '2. Project Objectives')
    obj_data = [
        ["1", "Facilitate communication and coordination within associations and circles."],
        ["2", "Provide a centralized platform for managing members, meetings, and events."],
        ["3", "Offer real-time messaging capabilities for groups and individuals."],
        ["4", "Streamline the creation and approval process of community circles and posts."]
    ]
    add_table_data(doc, obj_data, ["#", "Objective"])

    # 3. Technologies Used
    add_heading(doc, '3. Technologies Used')
    tech_data = [
        ["Frontend", "Flutter (Dart), Google Fonts, FontAwesome, Material UI"],
        ["Backend", "Node.js, Express.js"],
        ["Database", "MySQL (with mysql2 package)"],
        ["Real-time", "Socket.io"],
        ["Authentication", "JWT, Google Auth Library, Facebook Auth, Bcrypt"],
        ["File Management", "Multer (Local file uploads)"],
        ["Mail Service", "Nodemailer"]
    ]
    add_table_data(doc, tech_data, ["Domain", "Technology / Stack"])

    # 4. System Architecture
    add_heading(doc, '4. System Architecture')
    arch_data = [
        ["Client App", "Flutter mobile app communicating via RESTful APIs and WebSockets."],
        ["API Gateway", "Express.js routing requests to controllers and applying middleware (auth, logging)."],
        ["Business Logic", "Node.js controllers managing users, circles, posts, and real-time messaging."],
        ["Data Layer", "MySQL database storing relational data using explicit schemas."],
        ["Real-time Layer", "Socket.io facilitating live chat and notifications."]
    ]
    add_table_data(doc, arch_data, ["Component", "Role"])

    # 5. Frontend Structure
    add_heading(doc, '5. Frontend Structure')
    fe_data = [
        ["State Management", "Local State & Provider/Custom State Management"],
        ["Routing", "Navigator 2.0 / Standard Flutter Routing"],
        ["HTTP Client", "http package for API requests"],
        ["Local Storage", "shared_preferences & flutter_secure_storage"],
        ["Assets", "Local images (logo.png, bantou.jpg) and fonts"]
    ]
    add_table_data(doc, fe_data, ["Aspect", "Implementation"])

    # 6. Backend Structure
    add_heading(doc, '6. Backend Structure')
    be_data = [
        ["Entry Point", "server.js (Initialization & Middleware config)"],
        ["Routes", "Express routers (authRoutes, circleRoutes, messageRoutes, etc.)"],
        ["Controllers", "Business logic (authController.js, etc.)"],
        ["Models", "Database schemas and queries (User.js, Circle.js, Post.js, etc.)"],
        ["Middleware", "Custom middlewares for JWT validation and role checking"]
    ]
    add_table_data(doc, be_data, ["Layer", "Description"])

    # 7. Database Structure
    add_heading(doc, '7. Database Structure')
    db_data = [
        ["Users", "Stores user details, authentication credentials, and roles."],
        ["Associations", "Stores organization/company details."],
        ["Circles", "Stores group/circle information, schedules, and association links."],
        ["CircleAccessRequests", "Tracks join requests and their approval statuses."],
        ["Posts", "User and circle generated content, feed posts."],
        ["Messages", "Real-time chat histories between users or within circles."]
    ]
    add_table_data(doc, db_data, ["Table / Entity", "Purpose"])

    # 8. Main Features & Functionalities
    add_heading(doc, '8. Main Features & Functionalities')
    features_data = [
        ["Authentication", "Login (Local, Google, Facebook), Registration, Password Recovery."],
        ["Profile Management", "Edit personal info, professional details, and profile pictures."],
        ["Circle Management", "Create, join, and manage circles with specific scheduling (recurrences)."],
        ["Messaging", "Real-time 1-on-1 and Group chats with Socket.io."],
        ["Feed & Posts", "Create posts, view feeds, interact with community content."]
    ]
    add_table_data(doc, features_data, ["Feature Module", "Description"])

    # 9. User Roles and Permissions
    add_heading(doc, '9. User Roles and Permissions')
    roles_data = [
        ["Admin", "Full access to system, manage associations, approve requests."],
        ["Circle Admin", "Manage specific circles, approve member requests, schedule meetings."],
        ["Standard User", "Join circles, post content, chat with connections."],
        ["Guest", "Limited access, requires approval for features."]
    ]
    add_table_data(doc, roles_data, ["Role", "Permissions"])

    # 10. APIs and Services Used
    add_heading(doc, '10. APIs and Services Used')
    apis_data = [
        ["Authentication", "Google Sign-in API, Facebook Graph API"],
        ["Email", "SMTP Service via Nodemailer"],
        ["File Storage", "Local server storage via Multer"],
        ["Real-time Comm", "Socket.io connections"]
    ]
    add_table_data(doc, apis_data, ["Service", "Usage"])

    # 11. Screens / Pages Description
    add_heading(doc, '11. Screens / Pages Description')
    screens_data = [
        ["Auth Module", "AuthScreen, ProfileSetupScreen, EditProfileScreen, ForgotPassword"],
        ["Circles Module", "CircleDashboard, CreateCircle, MemberCircles, PendingRequests"],
        ["Messaging Module", "MessagesScreen, ChatScreen, GroupInfoScreen, NewGroupScreen"],
        ["Home / Feed", "MainShellScreen, PostsScreen, UserProfileScreen"]
    ]
    add_table_data(doc, screens_data, ["Module", "Key Screens"])

    # 12. Security & Authentication
    add_heading(doc, '12. Security & Authentication')
    sec_data = [
        ["Passwords", "Hashed using Bcrypt before DB insertion."],
        ["Tokens", "JWT for session management and route protection."],
        ["Storage", "Flutter Secure Storage for storing sensitive keys on device."],
        ["CORS", "Configured to restrict unauthorized cross-origin requests."]
    ]
    add_table_data(doc, sec_data, ["Security Measure", "Implementation"])

    # 13. Folder / File Organization
    add_heading(doc, '13. Folder / File Organization')
    folder_data = [
        ["/backend/models", "Database interaction logic"],
        ["/backend/controllers", "API endpoint handlers"],
        ["/backend/routes", "URL routing configurations"],
        ["/circleback/lib/screens", "UI Views organized by module (auth, circles, messaging)"],
        ["/circleback/lib/models", "Dart data classes for JSON serialization"]
    ]
    add_table_data(doc, folder_data, ["Directory", "Content"])

    # 14. Future Improvements
    add_heading(doc, '14. Future Improvements')
    future_data = [
        ["Notifications", "Implement Firebase Cloud Messaging (FCM) for Push Notifications."],
        ["Cloud Storage", "Migrate Multer local uploads to AWS S3 or Firebase Storage."],
        ["Media Sharing", "Enhance chat with multimedia support (video, voice notes)."],
        ["Web Version", "Compile the Flutter app for Web for broader accessibility."]
    ]
    add_table_data(doc, future_data, ["Improvement", "Description"])

    # Save Document
    doc.save('Bantou_Fiche_Technique.docx')
    print("Document successfully created!")

if __name__ == "__main__":
    create_document()
