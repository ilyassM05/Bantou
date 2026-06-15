import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def add_heading(doc, text, level=1):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0, 51, 102) # Dark Blue

def add_table_data(doc, data, headers):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Light Shading Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Add headers
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        for paragraph in hdr_cells[i].paragraphs:
            for run in paragraph.runs:
                run.bold = True
                
    # Add rows
    for row_data in data:
        row_cells = table.add_row().cells
        for i, cell_data in enumerate(row_data):
            row_cells[i].text = str(cell_data)
            
    doc.add_paragraph("")

def create_document():
    doc = Document()
    
    # Cover Page
    doc.add_paragraph('\n'*5)
    title = doc.add_paragraph('Suivi du Projet de Fin d’Études (PFE)')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.size = Pt(28)
        run.bold = True
        run.font.color.rgb = RGBColor(0, 51, 102)
        
    subtitle = doc.add_paragraph('Application de Gestion de Communauté : Bantou')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in subtitle.runs:
        run.font.size = Pt(18)
        run.italic = True
        
    doc.add_page_break()

    # 1. Architecture technique du projet
    add_heading(doc, '1. Architecture technique du projet')
    
    doc.add_heading('Présentation globale de l’architecture', level=2)
    if os.path.exists('architecture_schema.png'):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture('architecture_schema.png', width=Inches(6.0))
    else:
        p = doc.add_paragraph()
        run = p.add_run("[ Espace réservé : Insérer le schéma (image) de l'architecture ici ]")
        run.font.italic = True
        run.font.color.rgb = RGBColor(128, 128, 128) # Gris
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("")
    
    doc.add_heading('Les technologies utilisées', level=2)
    tech_data = [
        ["Couche", "Technologies"],
        ["Frontend (Client Mobile)", "Flutter (Dart), Material Design, Provider / State Management"],
        ["Backend (API & Serveur)", "Node.js, Express.js"],
        ["Base de données", "MySQL (Package: mysql2)"],
        ["Temps Réel", "Socket.io"],
        ["Sécurité & Authentification", "JWT (JSON Web Tokens), Bcrypt, Google Auth Library"],
        ["Gestion de fichiers", "Multer (stockage local)"]
    ]
    add_table_data(doc, tech_data[1:], tech_data[0])
    
    doc.add_heading('Organisation des différentes couches', level=2)
    couches_data = [
        ["Couche", "Rôle et Organisation"],
        ["Présentation (Frontend)", "Interface utilisateur divisée en modules (Auth, Circles, Messaging, Feed). Gère l'affichage et les interactions locales."],
        ["Métier (Backend - Controllers)", "Contient la logique métier de l'application, valide les requêtes, et applique les règles de sécurité (ex: authController, circleController)."],
        ["Services & Routage (Backend)", "Points d'entrée de l'API (Routes) et middlewares (Vérification JWT, upload de fichiers)."],
        ["Accès aux données (Models)", "Interaction directe avec la base MySQL via des requêtes SQL structurées (User.js, Circle.js)."]
    ]
    add_table_data(doc, couches_data[1:], couches_data[0])
    
    doc.add_heading('Les choix techniques et leurs justifications', level=2)
    choix_data = [
        ["Choix", "Justification"],
        ["Flutter", "Développement multiplateforme (iOS/Android) avec une seule base de code et d'excellentes performances natives."],
        ["Node.js & Express", "Performances élevées pour les I/O asynchrones, écosystème riche (NPM), et rapidité de mise en place des API REST."],
        ["MySQL", "Structure relationnelle robuste adaptée à la gestion des entités fortement liées (Utilisateurs, Cercles, Associations)."],
        ["Socket.io", "Nécessaire pour garantir une messagerie instantanée fluide et fiable avec gestion des déconnexions."]
    ]
    add_table_data(doc, choix_data[1:], choix_data[0])

    doc.add_page_break()

    # 2. Description détaillée de la base de données
    add_heading(doc, '2. Description détaillée de la base de données')
    
    doc.add_heading('Type de base de données utilisée', level=2)
    doc.add_paragraph("Le projet utilise une base de données relationnelle SQL (MySQL).")
    
    doc.add_heading('Le modèle de données et description des tables', level=2)
    tables_data = [
        ["Table", "Description & Attributs Principaux", "Relations"],
        ["Users", "Stocke les infos utilisateurs (id, name, email, password_hash, role).", "Liée à Circles, Posts, Messages (1:N)"],
        ["Associations", "Détails des organisations gérant les cercles (id, name, description).", "Liée à Circles (1:N)"],
        ["Circles", "Groupes communautaires (id, name, association_id, description, rules).", "Liée à Associations (N:1), Users (N:M)"],
        ["CircleAccessRequests", "Demandes d'adhésion (id, user_id, circle_id, status).", "Liée à Users (N:1), Circles (N:1)"],
        ["Posts", "Publications des utilisateurs ou cercles (id, author_id, content).", "Liée à Users, Circles"],
        ["Messages", "Historique du chat (id, sender_id, receiver_id/circle_id, text).", "Liée à Users, Circles"]
    ]
    add_table_data(doc, tables_data[1:], tables_data[0])
    
    doc.add_heading('Les règles de gestion (Données)', level=2)
    regles_data = [
        ["Règle", "Description"],
        ["RG_DB_01", "Un e-mail utilisateur doit être unique dans la table Users."],
        ["RG_DB_02", "Lorsqu'un cercle est supprimé, les demandes d'accès liées doivent être supprimées (Cascade)."],
        ["RG_DB_03", "Les mots de passe ne sont jamais stockés en clair (hash Bcrypt obligatoire)."],
        ["RG_DB_04", "Une demande d'accès ne peut être validée que par l'administrateur du cercle ou de l'association."]
    ]
    add_table_data(doc, regles_data[1:], regles_data[0])

    doc.add_page_break()

    # 3. Planning du projet
    add_heading(doc, '3. Planning du projet')
    doc.add_paragraph("Le tableau ci-dessous présente le découpage en sprints avec les User Stories et les règles métiers associées.")
    
    sprint_data = [
        ["Sprint", "User Story (US)", "Règles métiers"],
        
        # Sprint 1 : Environnement
        ["Sprint 1", "Mise en place de l'environnement de développement :\n- Configuration du dépôt Git.\n- Préparation de l'environnement de développement.\n- Mise en place des outils de gestion de projet et de collaboration.", "RM1.1: Tous les outils doivent être configurés et accessibles à l'équipe."],
        
        # Sprint 2 : Maquette
        ["Sprint 2", "Travail sur le design de l'application :\n- Création d'une maquette de l'application.", "RM2.1: La maquette doit être validée avant de commencer le développement UI."],
        
        # Sprint 3 (Ancien Sprint 1)
        ["Sprint 3", "En tant qu'utilisateur, je veux pouvoir m'inscrire et me connecter via Email/Mot de passe ou Google/Facebook.", "RM3.1: Le mot de passe doit contenir min. 8 caractères.\nRM3.2: L'email doit être unique.\nRM3.3: Un token JWT est généré avec une validité définie après connexion."],
        ["Sprint 3", "En tant qu'utilisateur, je veux pouvoir modifier mon profil (informations personnelles et professionnelles).", "RM3.4: La mise à jour de l'email nécessite une revalidation (optionnelle).\nRM3.5: Les champs obligatoires ne peuvent être vides."],
        
        # Sprint 4 (Ancien Sprint 2)
        ["Sprint 4", "En tant qu'administrateur, je veux créer et configurer une Association.", "RM4.1: Un utilisateur classique ne peut pas créer une association (rôle requis).\nRM4.2: Le nom de l'association est obligatoire."],
        ["Sprint 4", "En tant que responsable, je veux créer un Cercle au sein de l'Association et définir sa récurrence (réunions).", "RM4.3: La création d'un cercle requiert un ID d'association valide.\nRM4.4: Les règles de récurrence sont stockées au format JSON valide."],
        
        # Sprint 5 (Ancien Sprint 3)
        ["Sprint 5", "En tant qu'utilisateur, je veux rechercher des cercles et envoyer une demande d'adhésion.", "RM5.1: Un utilisateur ne peut envoyer qu'une seule demande en attente par cercle.\nRM5.2: L'historique d'appartenance empêche les doublons."],
        ["Sprint 5", "En tant qu'administrateur de cercle, je veux approuver ou rejeter les demandes d'adhésion.", "RM5.3: L'approbation change le statut en 'approved' et lie l'utilisateur au cercle.\nRM5.4: Une notification est envoyée à l'utilisateur concerné."],
        
        # Sprint 6 (Ancien Sprint 4)
        ["Sprint 6", "En tant qu'utilisateur, je veux publier un Post visible par ma communauté.", "RM6.1: Un post doit contenir du texte ou un média.\nRM6.2: Seuls les membres autorisés du cercle peuvent publier dans l'espace du cercle."],
        ["Sprint 6", "En tant qu'utilisateur, je veux envoyer des messages en temps réel dans un chat de groupe (Cercle).", "RM6.3: La connexion Socket n'est établie que pour les membres du cercle.\nRM6.4: Les messages sont persistés en base de données avant diffusion."]
    ]
    add_table_data(doc, sprint_data[1:], sprint_data[0])

    # Save Document
    doc.save('Evaluation_PFE_Bantou_v2.docx')
    print("Document successfully created! (v2)")

if __name__ == "__main__":
    create_document()
